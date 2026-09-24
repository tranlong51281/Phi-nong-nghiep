import streamlit as st
import pdfplumber
import openpyxl
from docx import Document
import re
import io

st.set_page_config(page_title="Phần mềm Xử Lý Hồ Sơ Đất Đai", page_icon="📝")

st.title("📝 Tự động hóa: PDF ➔ Excel & Word")
st.write("Dành cho hồ sơ Xác định nghĩa vụ tài chính / Chuyển mục đích sử dụng đất phường Bình Khê.")

# --- HÀM XỬ LÝ LÕI ---
def extract_info_from_pdf(pdf_file):
    text = ""
    with pdfplumber.open(pdf_file) as pdf:
        for page in pdf.pages:
            text += page.extract_text() + "\n"
    
    data = {
        "ten": "Không tìm thấy",
        "cccd": "Không tìm thấy",
        "thua_dat": "Không tìm thấy",
        "to_ban_do": "Không tìm thấy",
        "dien_tich_chuyen": "Không tìm thấy"
    }
    
    match_ten = re.search(r'Tên:\s*(Ông|Bà)?\s*([A-ZÀ-Ỹa-zà-ỹ\s]+)', text)
    if match_ten:
        data["ten"] = match_ten.group(2).strip()
        
    match_cccd = re.search(r'CCCD\s*(\d+)', text)
    if match_cccd:
        data["cccd"] = match_cccd.group(1)
        
    match_thua = re.search(r'Thửa đất số:\s*(\d+)', text)
    if match_thua:
        data["thua_dat"] = match_thua.group(1)
        
    match_bando = re.search(r'Tờ bản đồ số:.*?(\d+)', text)
    if match_bando:
        data["to_ban_do"] = match_bando.group(1)
        
    match_dt = re.search(r'Diện tích chuyển mục đích sử dụng đất:.*?([\d\,]+)', text)
    if match_dt:
        data["dien_tich_chuyen"] = match_dt.group(1)
        
    return data

# --- GIAO DIỆN WEB ---
# Thêm accept_multiple_files=True để cho phép bôi đen chọn nhiều file PDF
uploaded_pdfs = st.file_uploader("1. Tải lên CÁC file PDF hồ sơ", type=["pdf"], accept_multiple_files=True)
uploaded_excel = st.file_uploader("2. Tải lên file Excel mẫu (đuôi .xlsx)", type=["xlsx"])
uploaded_word = st.file_uploader("3. Tải lên file Word mẫu (đuôi .docx)", type=["docx"])

# Kiểm tra xem người dùng đã tải đủ file chưa (uploaded_pdfs giờ là 1 danh sách)
if uploaded_pdfs and uploaded_excel and uploaded_word:
    if st.button("🚀 BẮT ĐẦU XỬ LÝ HÀNG LOẠT"):
        with st.spinner(f"Đang xử lý {len(uploaded_pdfs)} hồ sơ..."):
            
            # Mở Excel mẫu 1 lần duy nhất để điền nhiều dòng
            wb = openpyxl.load_workbook(uploaded_excel)
            sheet = wb.active
            next_row = sheet.max_row + 1
            
            # Danh sách chứa các file Word kết quả
            word_results = []
            
            st.write("### 📊 Dữ liệu trích xuất được:")
            
            # Vòng lặp: Chạy qua TỪNG FILE PDF mà bạn tải lên
            for pdf_file in uploaded_pdfs:
                extracted_data = extract_info_from_pdf(pdf_file)
                st.json(extracted_data) # In ra màn hình để kiểm tra
                
                # Điền vào Excel (Mỗi file PDF điền thành 1 dòng mới)
                sheet.cell(row=next_row, column=1).value = extracted_data["ten"]
                sheet.cell(row=next_row, column=2).value = extracted_data["cccd"]
                sheet.cell(row=next_row, column=3).value = extracted_data["thua_dat"]
                sheet.cell(row=next_row, column=4).value = extracted_data["to_ban_do"]
                next_row += 1 # Cộng thêm 1 dòng cho hồ sơ tiếp theo
                
                # Tạo file Word riêng cho từng hồ sơ
                doc = Document(uploaded_word)
                for paragraph in doc.paragraphs:
                    if 'Phùng Văn Cầu' in paragraph.text:
                        paragraph.text = paragraph.text.replace('Phùng Văn Cầu', extracted_data["ten"])
                
                word_output = io.BytesIO()
                doc.save(word_output)
                word_output.seek(0)
                
                # Lưu file Word tạm vào bộ nhớ để lát nữa tạo nút tải
                word_results.append((extracted_data["ten"], word_output))
            
            # Lưu lại toàn bộ file Excel tổng hợp
            excel_output = io.BytesIO()
            wb.save(excel_output)
            excel_output.seek(0)
            
            st.success(f"✅ Đã xử lý xong toàn bộ {len(uploaded_pdfs)} hồ sơ!")
            st.write("---")
            st.write("### ⬇️ TẢI XUỐNG KẾT QUẢ")
            
            # 1 nút tải file Excel chứa tất cả các dòng
            st.download_button(
                label="📊 Tải file Excel (Tổng hợp)", 
                data=excel_output, 
                file_name="DuLieu_TongHop.xlsx", 
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            
            st.write("**Tải Thông báo Word của từng người:**")
            
            # In ra nhiều nút tải file Word (mỗi người 1 nút)
            for ten, word_out in word_results:
                st.download_button(
                    label=f"📝 Thông Báo - {ten}", 
                    data=word_out, 
                    file_name=f"ThongBao_{ten}.docx", 
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )